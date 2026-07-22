from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_SUBJECTS = ("语文", "数学", "英语", "物理", "化学", "政治", "历史", "地理", "生物")


class ScoreExcelImportError(ValueError):
    """Raised when an exam workbook is not in a usable format."""


@dataclass(frozen=True)
class ParsedScoreRow:
    row_number: int
    name: str
    student_no: str
    class_name: str
    scores: dict[str, float]
    invalid_scores: dict[str, str]


@dataclass(frozen=True)
class ParsedScoreWorkbook:
    rows: list[ParsedScoreRow]
    subjects: list[str]
    full_scores: dict[str, float]


_STUDENT_HEADERS = {
    "姓名": "name",
    "学生姓名": "name",
    "学号": "student_no",
    "学生学号": "student_no",
    "班级": "class_name",
    "班级名称": "class_name",
}
_SUBJECT_ALIASES = {
    "语文": "语文",
    "数学": "数学",
    "英语": "英语",
    "物理": "物理",
    "化学": "化学",
    "政治": "政治",
    "道德与法治": "政治",
    "道法": "政治",
    "历史": "历史",
    "地理": "地理",
    "生物": "生物",
}
_NON_SCORE_VALUES = {"", "-", "--", "缺考", "缺席", "未考", "不参加", "N/A", "NA"}


def ensure_score_import_template(file_path: str | Path) -> Path:
    path = Path(file_path)
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "单次考试成绩"
    worksheet.append(["姓名", "学号", "班级", *DEFAULT_SUBJECTS])
    worksheet.append(["示例学生", "2026001", "初三1班", 112, 108, 115, 86, 78, 90, 88, 82, 85])
    _style_sheet(worksheet)
    workbook.save(path)
    workbook.close()
    return path


def read_score_import_file(file_path: str | Path) -> ParsedScoreWorkbook:
    path = Path(file_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ScoreExcelImportError("请选择 .xlsx 或 .xlsm 格式的 Excel 文件。")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ScoreExcelImportError("无法打开 Excel 文件，请确认文件未损坏且未加密。") from exc

    try:
        worksheet = workbook.active
        header_row, student_columns, subject_columns, full_scores = _find_header_row(worksheet)
        if "name" not in student_columns:
            raise ScoreExcelImportError("未识别到“姓名”列。")
        if not subject_columns:
            raise ScoreExcelImportError("未识别到任何学科列，请至少保留一门参加考试的学科。")

        rows: list[ParsedScoreRow] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            name = _value_at(values, student_columns.get("name"))
            student_no = _value_at(values, student_columns.get("student_no"))
            class_name = _value_at(values, student_columns.get("class_name"))
            scores: dict[str, float] = {}
            invalid_scores: dict[str, str] = {}
            for subject, column_index in subject_columns.items():
                raw_value = _value_at(values, column_index)
                score = _parse_score(raw_value)
                if score is None:
                    continue
                if isinstance(score, str):
                    invalid_scores[subject] = score
                else:
                    scores[subject] = score
            if not name and not student_no and not class_name and not scores and not invalid_scores:
                continue
            rows.append(
                ParsedScoreRow(
                    row_number=row_number,
                    name=name,
                    student_no=student_no,
                    class_name=class_name,
                    scores=scores,
                    invalid_scores=invalid_scores,
                )
            )
        return ParsedScoreWorkbook(rows=rows, subjects=list(subject_columns), full_scores=full_scores)
    finally:
        workbook.close()


def _find_header_row(worksheet) -> tuple[int, dict[str, int], dict[str, int], dict[str, float]]:
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True),
        start=1,
    ):
        student_columns: dict[str, int] = {}
        subject_columns: dict[str, int] = {}
        full_scores: dict[str, float] = {}
        for column_index, value in enumerate(values):
            header = _normalize_header(_text_value(value))
            field = _STUDENT_HEADERS.get(header)
            if field and field not in student_columns:
                student_columns[field] = column_index
                continue
            subject, full_score = _subject_from_header(header)
            if subject and subject not in subject_columns:
                subject_columns[subject] = column_index
                if full_score is not None:
                    full_scores[subject] = full_score
        if "name" in student_columns and subject_columns:
            return row_number, student_columns, subject_columns, full_scores
    raise ScoreExcelImportError("未找到成绩表头，请使用系统提供的成绩导入模板。")


def _subject_from_header(header: str) -> tuple[str | None, float | None]:
    cleaned = header.replace("成绩", "").replace("分数", "")
    for alias, subject in _SUBJECT_ALIASES.items():
        if cleaned == alias:
            return subject, _header_full_score(header)
        if cleaned.startswith(alias) and cleaned[len(alias) :].startswith(("(", "（", "[", "【")):
            return subject, _header_full_score(header)
    return None, None


def _header_full_score(header: str) -> float | None:
    match = re.search(r"[（(\[【]\s*(\d+(?:\.\d+)?)", header)
    return float(match.group(1)) if match else None


def _parse_score(value: str) -> float | str | None:
    normalized = value.upper().strip()
    if normalized in _NON_SCORE_VALUES:
        return None
    try:
        score = float(normalized)
    except ValueError:
        return value
    if score < 0:
        return value
    return score


def _value_at(values: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(values):
        return ""
    return _text_value(values[index])


def _text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: str) -> str:
    return value.replace(" ", "").replace("\n", "").replace("\r", "").strip()


def _style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="00A870")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index in range(1, worksheet.max_column + 1):
        values = [str(worksheet.cell(row, column_index).value or "") for row in range(1, worksheet.max_row + 1)]
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max(map(len, values), default=10) + 3, 11), 18
        )
