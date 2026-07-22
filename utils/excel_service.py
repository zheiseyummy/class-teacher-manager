from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.quality_scoring import SEMESTER_RULES


STUDENT_IMPORT_HEADERS = [
    "姓名",
    "性别",
    "学号",
    "班级",
    "年级",
    "座号",
    "出生日期",
    "身份证号",
    "民族",
    "学生备注",
    "家长姓名",
    "与学生关系",
    "联系电话",
    "微信号",
    "工作单位",
    "家庭地址",
    "家长备注",
]

HEADER_ALIASES = {
    "姓名": "姓名",
    "学生姓名": "姓名",
    "性别": "性别",
    "学号": "学号",
    "学生学号": "学号",
    "班级": "班级",
    "班级名称": "班级",
    "年级": "年级",
    "座号": "座号",
    "座位号": "座号",
    "出生日期": "出生日期",
    "生日": "出生日期",
    "身份证号": "身份证号",
    "身份证号码": "身份证号",
    "民族": "民族",
    "学生备注": "学生备注",
    "备注": "学生备注",
    "家长姓名": "家长姓名",
    "联系人姓名": "家长姓名",
    "与学生关系": "与学生关系",
    "关系": "与学生关系",
    "联系电话": "联系电话",
    "家长电话": "联系电话",
    "联系人电话": "联系电话",
    "电话": "联系电话",
    "手机号": "联系电话",
    "微信号": "微信号",
    "微信": "微信号",
    "工作单位": "工作单位",
    "家庭地址": "家庭地址",
    "地址": "家庭地址",
    "家长备注": "家长备注",
    "联系人备注": "家长备注",
}


class ExcelImportError(ValueError):
    """Raised when an Excel workbook cannot be used as a student import file."""


@dataclass(frozen=True)
class ParsedStudentRow:
    row_number: int
    student: dict[str, Any]
    guardian: dict[str, Any]


def ensure_student_import_template(template_path: Path) -> Path:
    """Create the standard student import template once, without overwriting user files."""
    if template_path.exists():
        return template_path
    template_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学生信息"
    worksheet.append(STUDENT_IMPORT_HEADERS)
    _style_sheet(worksheet)
    for index, header in enumerate(STUDENT_IMPORT_HEADERS, start=1):
        column = worksheet.column_dimensions[get_column_letter(index)]
        column.width = max(12, min(22, len(header) * 3 + 6))
        if header in {"学号", "身份证号", "联系电话", "微信号"}:
            column.number_format = "@"
    worksheet.freeze_panes = "A2"
    workbook.save(template_path)
    return template_path


def read_student_import_file(file_path: str | Path) -> list[ParsedStudentRow]:
    path = Path(file_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelImportError("请选择 .xlsx 或 .xlsm 格式的 Excel 文件。")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ExcelImportError("无法打开 Excel 文件，请确认文件未损坏且未被加密。") from exc

    try:
        worksheet = workbook.active
        header_row, columns = _find_header_row(worksheet)
        missing = [header for header in ("姓名", "班级") if header not in columns]
        if missing:
            raise ExcelImportError(f"未识别到必填列：{'、'.join(missing)}。")

        rows: list[ParsedStudentRow] = []
        for row_number, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            row_values = {
                header: _text_value(values[column_index]) if column_index < len(values) else ""
                for header, column_index in columns.items()
            }
            if not any(row_values.values()):
                continue
            rows.append(
                ParsedStudentRow(
                    row_number=row_number,
                    student={
                        "name": row_values.get("姓名", ""),
                        "gender": row_values.get("性别", ""),
                        "student_no": row_values.get("学号", ""),
                        "class_name": row_values.get("班级", ""),
                        "grade": row_values.get("年级", ""),
                        "seat_no": row_values.get("座号", ""),
                        "birth_date": _parse_date(row_values.get("出生日期", "")),
                        "id_card": row_values.get("身份证号", ""),
                        "ethnicity": row_values.get("民族", ""),
                        "note": row_values.get("学生备注", ""),
                    },
                    guardian={
                        "name": row_values.get("家长姓名", ""),
                        "relationship": row_values.get("与学生关系", ""),
                        "phone": row_values.get("联系电话", ""),
                        "wechat": row_values.get("微信号", ""),
                        "workplace": row_values.get("工作单位", ""),
                        "address": row_values.get("家庭地址", ""),
                        "note": row_values.get("家长备注", ""),
                    },
                )
            )
        return rows
    finally:
        workbook.close()


def write_student_export(file_path: str | Path, students: Iterable[dict[str, Any]], title: str = "学生信息") -> Path:
    headers = [
        "姓名",
        "性别",
        "学号",
        "班级",
        "年级",
        "座号",
        "出生日期",
        "身份证号",
        "民族",
        "家长姓名",
        "与学生关系",
        "联系电话",
        "微信号",
        "工作单位",
        "家庭地址",
        "学生备注",
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title[:31]
    worksheet.append(headers)
    for student in students:
        guardians = student.get("guardians", [])
        worksheet.append(
            [
                student.get("name", ""),
                student.get("gender", ""),
                student.get("student_no", ""),
                student.get("class_name", ""),
                student.get("grade", ""),
                student.get("seat_no", ""),
                student.get("birth_date", ""),
                student.get("id_card", ""),
                student.get("ethnicity", ""),
                _join_guardian_values(guardians, "name"),
                _join_guardian_values(guardians, "relationship"),
                _join_guardian_values(guardians, "phone"),
                _join_guardian_values(guardians, "wechat"),
                _join_guardian_values(guardians, "workplace"),
                _join_guardian_values(guardians, "address"),
                student.get("note", ""),
            ]
        )
    _style_sheet(worksheet)
    return _save_workbook(workbook, file_path)


def write_guardian_directory(file_path: str | Path, students: Iterable[dict[str, Any]]) -> Path:
    headers = ["学生姓名", "班级", "学号", "家长姓名", "关系", "联系电话", "微信号", "工作单位", "家庭地址", "备注"]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "家长通讯录"
    worksheet.append(headers)
    for student in students:
        for guardian in student.get("guardians", []):
            worksheet.append(
                [
                    student.get("name", ""),
                    student.get("class_name", ""),
                    student.get("student_no", ""),
                    guardian.get("name", ""),
                    guardian.get("relationship", ""),
                    guardian.get("phone", ""),
                    guardian.get("wechat", ""),
                    guardian.get("workplace", ""),
                    guardian.get("address", ""),
                    guardian.get("note", ""),
                ]
            )
    _style_sheet(worksheet)
    return _save_workbook(workbook, file_path)


def write_attendance_export(file_path: str | Path, records: Iterable[dict[str, Any]]) -> Path:
    """Write leave and attendance records using the application's standard Excel styling."""

    headers = ["日期时间", "结束时间", "学生", "学号", "班级", "类型", "状态", "时长", "事由", "备注"]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "请假与考勤"
    worksheet.append(headers)
    for record in records:
        worksheet.append(
            [
                record.get("start_time_display", ""),
                record.get("end_time_display", ""),
                record.get("student_name", ""),
                record.get("student_no", ""),
                record.get("class_name", ""),
                record.get("record_type", ""),
                record.get("approval_status", ""),
                record.get("duration_display", ""),
                record.get("reason", ""),
                record.get("note", ""),
            ]
        )
    _style_sheet(worksheet)
    return _save_workbook(workbook, file_path)


def write_moral_export(file_path: str | Path, records: Iterable[dict[str, Any]]) -> Path:
    """Write moral-development activities and honors to a formatted Excel workbook."""

    headers = ["日期", "学生", "学号", "班级", "类别", "活动或奖项", "级别", "主办方", "德育积分", "备注"]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "德育评价"
    worksheet.append(headers)
    for record in records:
        worksheet.append(
            [
                record.get("record_date_display", ""),
                record.get("student_name", ""),
                record.get("student_no", ""),
                record.get("class_name", ""),
                record.get("category", ""),
                record.get("title", ""),
                record.get("award_level", ""),
                record.get("organizer", ""),
                record.get("points", 0),
                record.get("note", ""),
            ]
        )
    _style_sheet(worksheet)
    return _save_workbook(workbook, file_path)


def write_quality_export(
    file_path: str | Path,
    students: Iterable[dict[str, Any]],
    dimensions: list[str],
) -> Path:
    student_rows = list(students)
    workbook = Workbook()
    total_sheet = workbook.active
    total_sheet.title = "综合素质总表"

    base_headers = ["姓名", "学号", "班级"]
    rating_headers = [f"{rule.label}-{dimension}" for rule in SEMESTER_RULES for dimension in dimensions]
    dimension_headers = [f"{dimension}得分" for dimension in dimensions]
    dimension_rank_headers = [f"{dimension}班级排名" for dimension in dimensions]
    dimension_level_headers = [f"{dimension}最终等级" for dimension in dimensions]
    total_sheet.append(
        base_headers
        + rating_headers
        + dimension_headers
        + dimension_rank_headers
        + dimension_level_headers
        + ["当前累计", "最终总分", "五维总排名", "九下名单", "完成状态", "最终评定状态"]
    )
    for student in student_rows:
        summary = student["summary"]
        ratings = student["ratings"]
        final = student.get("final", {})
        total_sheet.append(
            _quality_base_values(student)
            + [ratings.get((rule.key, dimension), "") for rule in SEMESTER_RULES for dimension in dimensions]
            + [summary["dimension_scores"].get(dimension, 0) for dimension in dimensions]
            + [final.get("dimension_ranks", {}).get(dimension, "") for dimension in dimensions]
            + [final.get("final_levels", {}).get(dimension, "") for dimension in dimensions]
            + [
                summary["total_score"],
                summary["total_score"] if final.get("in_final_roster") else "",
                final.get("total_rank") or "",
                "是" if final.get("in_final_roster") else "否",
                _quality_completion_status(summary),
                _quality_final_status(final),
            ]
        )
    _style_sheet(total_sheet)
    _highlight_quality_gaps(total_sheet)

    dimension_sheet = workbook.create_sheet("五维得分")
    dimension_sheet.append(
        base_headers
        + dimension_headers
        + dimension_rank_headers
        + dimension_level_headers
        + ["五维总分", "五维总排名", "完成状态", "最终评定状态"]
    )
    for student in student_rows:
        summary = student["summary"]
        final = student.get("final", {})
        dimension_sheet.append(
            _quality_base_values(student)
            + [summary["dimension_scores"].get(dimension, 0) for dimension in dimensions]
            + [final.get("dimension_ranks", {}).get(dimension, "") for dimension in dimensions]
            + [final.get("final_levels", {}).get(dimension, "") for dimension in dimensions]
            + [
                summary["total_score"] if final.get("in_final_roster") else "",
                final.get("total_rank") or "",
                _quality_completion_status(summary),
                _quality_final_status(final),
            ]
        )
    _style_sheet(dimension_sheet)
    _highlight_quality_gaps(dimension_sheet)

    semester_sheet = workbook.create_sheet("学期得分")
    semester_headers = [f"{rule.label}得分" for rule in SEMESTER_RULES]
    semester_sheet.append(base_headers + semester_headers + ["当前累计", "完成学期数", "完成状态"])
    for student in student_rows:
        summary = student["summary"]
        semester_sheet.append(
            _quality_base_values(student)
            + [summary["semester_scores"].get(rule.key, 0) for rule in SEMESTER_RULES]
            + [
                summary["total_score"],
                summary["completed_terms"],
                _quality_completion_status(summary),
            ]
        )
    _style_sheet(semester_sheet)

    final_sheet = workbook.create_sheet("最终评定")
    final_dimension_headers = [
        header
        for dimension in dimensions
        for header in (
            f"{dimension}累计得分",
            f"{dimension}班级排名",
            f"{dimension}系统等级",
            f"{dimension}最终等级",
        )
    ]
    final_sheet.append(
        base_headers
        + rating_headers
        + final_dimension_headers
        + ["五维总分", "五维总排名", "数据状态", "复核状态"]
    )
    for student in student_rows:
        final = student.get("final", {})
        if not final.get("in_final_roster"):
            continue
        summary = student["summary"]
        ratings = student["ratings"]
        final_dimension_values: list[Any] = []
        for dimension in dimensions:
            final_dimension_values.extend(
                [
                    summary["dimension_scores"].get(dimension, 0),
                    final.get("dimension_ranks", {}).get(dimension, ""),
                    final.get("automatic_levels", {}).get(dimension, ""),
                    final.get("final_levels", {}).get(dimension, ""),
                ]
            )
        final_sheet.append(
            _quality_base_values(student)
            + [ratings.get((rule.key, dimension), "N/A") for rule in SEMESTER_RULES for dimension in dimensions]
            + final_dimension_values
            + [
                summary["total_score"],
                final.get("total_rank") or "",
                "含 N/A" if final.get("contains_na") else "完整",
                _quality_final_status(final),
            ]
        )
    _style_sheet(final_sheet)
    _highlight_quality_gaps(final_sheet)
    return _save_workbook(workbook, file_path)


def _find_header_row(worksheet) -> tuple[int, dict[str, int]]:
    best_row = 0
    best_columns: dict[str, int] = {}
    for row_number, values in enumerate(worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True), start=1):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(values):
            header = _normalize_header(_text_value(value))
            canonical = HEADER_ALIASES.get(header)
            if canonical and canonical not in columns:
                columns[canonical] = column_index
        if len(columns) > len(best_columns):
            best_row = row_number
            best_columns = columns
    if len(best_columns) < 2:
        raise ExcelImportError("未找到学生信息表头，请使用系统提供的 Excel 模板。")
    return best_row, best_columns


def _parse_date(value: str) -> date | str | None:
    if not value:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            continue
    # Keep invalid values for the row-level validator so one bad date does not abort the whole import.
    return value


def _text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: str) -> str:
    return value.replace(" ", "").replace("\n", "").replace("\r", "")


def _join_guardian_values(guardians: list[dict[str, Any]], key: str) -> str:
    return "；".join(str(guardian.get(key) or "") for guardian in guardians if guardian.get(key))


def _quality_base_values(student: dict[str, Any]) -> list[str]:
    return [student.get("name", ""), student.get("student_no", ""), student.get("class_name", "")]


def _quality_completion_status(summary: dict[str, Any]) -> str:
    if not summary["final_ready"]:
        return "未完成"
    return "含 N/A" if summary.get("contains_na") else "已完成"


def _quality_final_status(final: dict[str, Any]) -> str:
    if not final.get("in_final_roster"):
        return "不在九下名单"
    if not final.get("generated"):
        return "待生成"
    return "已锁定" if final.get("is_locked") else "待复核"


def _highlight_quality_gaps(worksheet) -> None:
    warning_fill = PatternFill("solid", fgColor="FFF3CD")
    warning_font = Font(color="8A5A00", bold=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            value = str(cell.value or "")
            if value == "N/A" or "含 N/A" in value:
                cell.fill = warning_fill
                cell.font = warning_font


def _style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="00A870")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index in range(1, worksheet.max_column + 1):
        values = [str(worksheet.cell(row, column_index).value or "") for row in range(1, worksheet.max_row + 1)]
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max(map(len, values), default=10) + 3, 11), 28)


def _save_workbook(workbook: Workbook, file_path: str | Path) -> Path:
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
