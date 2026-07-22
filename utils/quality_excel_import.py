from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utils.quality_scoring import SEMESTER_RULES


class QualityExcelImportError(ValueError):
    """Raised when a quality-evaluation workbook cannot be read safely."""


@dataclass(frozen=True)
class ParsedQualityRow:
    semester_key: str
    sheet_name: str
    row_number: int
    name: str
    student_no: str
    ratings: dict[str, str]


@dataclass(frozen=True)
class ParsedQualityWorkbook:
    rows: list[ParsedQualityRow]
    semester_sheets: dict[str, str]


_SEMESTER_ALIASES = {
    "junior_1_1": ("七上", "七年级上", "初一上", "初一上学期"),
    "junior_1_2": ("七下", "七年级下", "初一下", "初一下学期"),
    "junior_2_1": ("八上", "八年级上", "初二上", "初二上学期"),
    "junior_2_2": ("八下", "八年级下", "初二下", "初二下学期"),
    "junior_3_1": ("九上", "九年级上", "初三上", "初三上学期"),
    "junior_3_2": ("九下", "九年级下", "初三下", "初三下学期"),
}


def _normalize_name(value: str) -> str:
    return value.replace(" ", "").replace("\n", "").replace("\r", "").strip()


_SHEET_TO_SEMESTER = {
    _normalize_name(alias): semester_key
    for semester_key, aliases in _SEMESTER_ALIASES.items()
    for alias in aliases
}


def read_quality_import_file(file_path: str | Path, dimensions: list[str]) -> ParsedQualityWorkbook:
    path = Path(file_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise QualityExcelImportError("请选择 .xlsx 或 .xlsm 格式的 Excel 文件。")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise QualityExcelImportError("无法打开 Excel 文件，请确认文件未损坏且未加密。") from exc

    try:
        semester_sheets = {
            _SHEET_TO_SEMESTER[_normalize_name(worksheet.title)]: worksheet.title
            for worksheet in workbook.worksheets
            if _normalize_name(worksheet.title) in _SHEET_TO_SEMESTER
        }
        missing_semesters = [rule.label for rule in SEMESTER_RULES if rule.key not in semester_sheets]
        if missing_semesters:
            raise QualityExcelImportError(
                f"未识别到以下学期工作表：{'、'.join(missing_semesters)}。"
                "工作表可命名为七上、七下、八上、八下、九上、九下。"
            )

        rows: list[ParsedQualityRow] = []
        for rule in SEMESTER_RULES:
            sheet_name = semester_sheets[rule.key]
            worksheet = workbook[sheet_name]
            header_row, columns = _find_header_row(worksheet, dimensions)
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                name = _text_value(values[columns["姓名"]]) if columns["姓名"] < len(values) else ""
                student_no_column = columns.get("学号")
                student_no = (
                    _text_value(values[student_no_column])
                    if student_no_column is not None and student_no_column < len(values)
                    else ""
                )
                ratings = {
                    dimension: _text_value(values[columns[dimension]])
                    if columns[dimension] < len(values)
                    else ""
                    for dimension in dimensions
                }
                if not name and not any(ratings.values()):
                    continue
                rows.append(
                    ParsedQualityRow(
                        semester_key=rule.key,
                        sheet_name=sheet_name,
                        row_number=row_number,
                        name=name,
                        student_no=student_no,
                        ratings=ratings,
                    )
                )
        return ParsedQualityWorkbook(rows=rows, semester_sheets=semester_sheets)
    finally:
        workbook.close()


def _find_header_row(worksheet, dimensions: list[str]) -> tuple[int, dict[str, int]]:
    required_headers = ("姓名", *dimensions)
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True),
        start=1,
    ):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(values):
            header = _normalize_name(_text_value(value))
            if header in {"学号", "学生学号", "学籍号"} and "学号" not in columns:
                columns["学号"] = column_index
            for required in required_headers:
                aliases = {_normalize_name(required)}
                if required == "实践与创新":
                    aliases.add("实践创新")
                if header in aliases and required not in columns:
                    columns[required] = column_index
        if all(header in columns for header in required_headers):
            return row_number, columns
    raise QualityExcelImportError(
        f"工作表“{worksheet.title}”未找到完整表头：姓名、{'、'.join(dimensions)}。"
    )


def _text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
