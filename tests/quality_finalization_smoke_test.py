from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def run() -> None:
    _test_allocation_rules()
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        config.DATA_DIR = root / "data"
        config.BACKUP_DIR = root / "backups"
        config.DATABASE_PATH = config.DATA_DIR / "class_manager.db"
        config.DATABASE_URL = f"sqlite:///{config.DATABASE_PATH.as_posix()}"

        from controllers.quality_controller import QualityController, QualityDataError
        from controllers.student_controller import StudentController
        from database.connection import engine
        from database.init_db import initialize_database
        from openpyxl import load_workbook
        from utils.quality_scoring import SEMESTER_RULES

        try:
            initialize_database()
            student_controller = StudentController()
            class_id = student_controller.create_class({"name": "初三测试班", "grade": "初三"})
            quality_controller = QualityController()
            dimensions = [item["name"] for item in quality_controller.list_dimensions()]

            candidate_ids: list[int] = []
            for index in range(1, 11):
                student_id = student_controller.create_student(
                    {
                        "class_id": class_id,
                        "name": f"学生{index:02d}",
                        "student_no": f"Q{index:03d}",
                    },
                    [],
                )
                candidate_ids.append(student_id)
                for term_index, rule in enumerate(SEMESTER_RULES):
                    if index == 10 and term_index < 3:
                        level = "N/A"
                    else:
                        level = ("A", "B", "C")[(index + term_index) % 3]
                    quality_controller.save_evaluation(
                        student_id,
                        rule.key,
                        {dimension: level for dimension in dimensions},
                    )

            archived_id = student_controller.create_student(
                {"class_id": class_id, "name": "九下前转出", "student_no": "Q999"},
                [],
            )
            for rule in SEMESTER_RULES[:-1]:
                quality_controller.save_evaluation(
                    archived_id,
                    rule.key,
                    {dimension: "A" for dimension in dimensions},
                )

            generated = quality_controller.generate_final_results(class_id)
            assert generated["student_count"] == 10
            assert generated["result_count"] == 50

            review = quality_controller.get_final_review(class_id, dimensions[0])
            assert review["generated"]
            assert review["candidate_count"] == 10
            assert len(review["rows"]) == 10
            assert review["incomplete_count"] == 1
            assert review["rows"][0]["student_id"] == candidate_ids[-1]
            assert all(row["student_id"] != archived_id for row in review["rows"])

            levels_by_score: dict[float, set[str]] = {}
            ranks_by_score: dict[float, set[int]] = {}
            for row in review["rows"]:
                levels_by_score.setdefault(row["score"], set()).add(row["automatic_level"])
                ranks_by_score.setdefault(row["score"], set()).add(row["rank"])
            assert all(len(levels) == 1 for levels in levels_by_score.values())
            assert all(len(ranks) == 1 for ranks in ranks_by_score.values())

            from PySide6.QtWidgets import QApplication
            from views.quality_final_review_dialog import QualityFinalReviewDialog

            app = QApplication.instance() or QApplication([])
            dialog = QualityFinalReviewDialog(quality_controller, class_id)
            assert dialog.table.rowCount() == 10
            assert dialog.dimension_box.count() == 5
            assert dialog.lock_button.isEnabled()
            dialog.close()
            app.processEvents()

            result = review["rows"][0]
            adjusted_level = "C" if result["automatic_level"] != "C" else "B"
            quality_controller.update_final_level(result["result_id"], adjusted_level)
            adjusted_review = quality_controller.get_final_review(class_id, dimensions[0])
            adjusted_row = next(row for row in adjusted_review["rows"] if row["result_id"] == result["result_id"])
            assert adjusted_row["final_level"] == adjusted_level
            assert adjusted_row["is_manually_adjusted"]

            quality_controller.set_finalization_locked(class_id, True)
            locked_review = quality_controller.get_final_review(class_id, dimensions[0])
            assert locked_review["is_locked"]
            try:
                quality_controller.update_final_level(result["result_id"], "B")
            except QualityDataError:
                pass
            else:
                raise AssertionError("Locked final levels must not be editable")
            try:
                quality_controller.save_evaluation(
                    candidate_ids[0],
                    SEMESTER_RULES[-1].key,
                    {dimension: "A" for dimension in dimensions},
                )
            except QualityDataError:
                pass
            else:
                raise AssertionError("Locked source ratings must not be editable")

            export_path = quality_controller.export_quality_to_excel(root / "最终评定档案.xlsx", class_id)
            workbook = load_workbook(export_path, read_only=True)
            try:
                assert workbook.sheetnames == ["综合素质总表", "五维得分", "学期得分", "最终评定"]
                final_sheet = workbook["最终评定"]
                headers = [cell.value for cell in next(final_sheet.iter_rows(max_row=1))]
                assert final_sheet.max_row == 11
                assert "思想品德班级排名" in headers
                assert "思想品德系统等级" in headers
                assert "思想品德最终等级" in headers
                assert "五维总排名" in headers
                assert "复核状态" in headers
            finally:
                workbook.close()

            quality_controller.set_finalization_locked(class_id, False)
            assert not quality_controller.get_final_review(class_id, dimensions[0])["is_locked"]
        finally:
            engine.dispose()


def _test_allocation_rules() -> None:
    from utils.quality_scoring import allocate_final_levels, competition_ranks

    ranks = competition_ranks({1: 10, 2: 9, 3: 9, 4: 8})
    assert ranks == {1: 1, 2: 2, 3: 2, 4: 4}

    unique_scores = {student_id: 100 - student_id for student_id in range(1, 51)}
    unique_levels = allocate_final_levels(unique_scores)
    assert list(unique_levels.values()).count("A") == 30
    assert list(unique_levels.values()).count("C") == 2

    tied_48_49 = {student_id: 1000 - student_id for student_id in range(1, 48)}
    tied_48_49.update({48: 10, 49: 10, 50: 9})
    tied_levels = allocate_final_levels(tied_48_49)
    assert tied_levels[48] == tied_levels[49] == "B"
    assert tied_levels[50] == "C"

    tied_49_50 = {student_id: 1000 - student_id for student_id in range(1, 49)}
    tied_49_50.update({49: 9, 50: 9})
    tied_bottom_levels = allocate_final_levels(tied_49_50)
    assert tied_bottom_levels[49] == tied_bottom_levels[50] == "C"

    a_boundary = {student_id: 1000 - student_id for student_id in range(1, 30)}
    a_boundary.update({30: 900, 31: 900})
    a_boundary.update({student_id: 800 - student_id for student_id in range(32, 51)})
    a_boundary_levels = allocate_final_levels(a_boundary)
    assert list(a_boundary_levels.values()).count("A") == 31
    assert a_boundary_levels[30] == a_boundary_levels[31] == "A"

    oversized_tie = {student_id: 1000 - student_id for student_id in range(1, 29)}
    oversized_tie.update({student_id: 900 for student_id in range(29, 33)})
    oversized_tie.update({student_id: 800 - student_id for student_id in range(33, 51)})
    oversized_levels = allocate_final_levels(oversized_tie)
    assert list(oversized_levels.values()).count("A") == 28
    assert {oversized_levels[student_id] for student_id in range(29, 33)} == {"B"}


if __name__ == "__main__":
    run()
    print("Quality finalization smoke test passed.")
