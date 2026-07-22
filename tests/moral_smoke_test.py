from __future__ import annotations

import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def run() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        config.DATA_DIR = root / "data"
        config.BACKUP_DIR = root / "backups"
        config.DATABASE_PATH = config.DATA_DIR / "class_manager.db"
        config.DATABASE_URL = f"sqlite:///{config.DATABASE_PATH.as_posix()}"

        from controllers.moral_controller import MoralController, MoralDataError
        from controllers.student_controller import StudentController
        from database.connection import engine
        from database.init_db import initialize_database
        from openpyxl import load_workbook

        try:
            initialize_database()
            students = StudentController()
            class_a = students.create_class({"name": "初三1班", "grade": "初三"})
            class_b = students.create_class({"name": "初三2班", "grade": "初三"})
            student_a = students.create_student(
                {"class_id": class_a, "name": "李明", "student_no": "001"}, []
            )
            student_b = students.create_student(
                {"class_id": class_b, "name": "张婷", "student_no": "002"}, []
            )

            controller = MoralController()
            activity_id = controller.create_record(
                {
                    "student_id": student_a,
                    "record_date": date(2026, 9, 5),
                    "category": "集体活动",
                    "title": "校园艺术节志愿服务",
                    "organizer": "学校团委",
                    "points": 2,
                    "note": "负责秩序维护",
                }
            )
            award_id = controller.create_record(
                {
                    "student_id": student_a,
                    "record_date": date(2026, 9, 20),
                    "category": "获奖荣誉",
                    "title": "三好学生",
                    "award_level": "校级",
                    "organizer": "学校",
                    "points": 5.5,
                }
            )
            controller.create_record(
                {
                    "student_id": student_b,
                    "record_date": date(2026, 9, 25),
                    "category": "志愿服务",
                    "title": "社区环保宣传",
                    "points": 3,
                }
            )

            class_a_rows = controller.list_records(class_id=class_a)
            assert len(class_a_rows) == 2
            assert class_a_rows[0]["id"] == award_id
            summary = controller.get_summary(class_id=class_a)
            assert summary == {"total": 2, "activities": 1, "awards": 1, "points": 7.5}
            student_summary = controller.get_student_summary(student_a)
            assert student_summary == {"total": 2, "points": 7.5, "activities": 1, "awards": 1}

            assert len(controller.list_records(category="获奖荣誉")) == 1
            assert len(controller.list_records(keyword="艺术节")) == 1
            assert len(
                controller.list_records(start_date=date(2026, 9, 20), end_date=date(2026, 9, 20))
            ) == 1

            controller.update_record(
                activity_id,
                {
                    "student_id": student_a,
                    "record_date": date(2026, 9, 6),
                    "category": "班级服务",
                    "title": "班级图书角整理",
                    "award_level": "",
                    "organizer": "初三1班",
                    "points": 2.5,
                    "note": "持续一周",
                },
            )
            updated = controller.get_record(activity_id)
            assert updated["category"] == "班级服务"
            assert updated["points_display"] == "2.5"
            assert controller.get_student_summary(student_a)["points"] == 8

            try:
                controller.create_record(
                    {
                        "student_id": student_a,
                        "record_date": date(2026, 9, 30),
                        "category": "获奖荣誉",
                        "title": "",
                    }
                )
            except MoralDataError:
                pass
            else:
                raise AssertionError("A record without a title should be rejected.")

            export_path = controller.export_records_to_excel(root / "德育评价.xlsx", class_id=class_a)
            workbook = load_workbook(export_path, read_only=True)
            worksheet = workbook["德育评价"]
            headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
            rows = list(worksheet.iter_rows(min_row=2, values_only=True))
            workbook.close()
            assert headers == ["日期", "学生", "学号", "班级", "类别", "活动或奖项", "级别", "主办方", "德育积分", "备注"]
            assert len(rows) == 2

            controller.delete_record(award_id)
            assert controller.get_student_summary(student_a)["total"] == 1
        finally:
            engine.dispose()


if __name__ == "__main__":
    run()
    print("Moral evaluation smoke test passed.")
