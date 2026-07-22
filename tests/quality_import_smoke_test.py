from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def run() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        config.DATA_DIR = root / "data"
        config.BACKUP_DIR = root / "backups"
        config.DATABASE_PATH = config.DATA_DIR / "class_manager.db"
        config.DATABASE_URL = f"sqlite:///{config.DATABASE_PATH.as_posix()}"

        from controllers.quality_controller import QualityController
        from controllers.student_controller import StudentController
        from database.connection import engine
        from database.init_db import initialize_database
        from openpyxl import Workbook, load_workbook
        from utils.quality_scoring import SEMESTER_RULES

        try:
            initialize_database()
            student_controller = StudentController()
            class_id = student_controller.create_class({"name": "初三1班", "grade": "初三"})
            complete_student_id = student_controller.create_student(
                {"class_id": class_id, "name": "完整学生", "student_no": "1001"}, []
            )
            transfer_student_id = student_controller.create_student(
                {"class_id": class_id, "name": "转入学生", "student_no": "1002"}, []
            )

            quality_controller = QualityController()
            dimensions = [item["name"] for item in quality_controller.list_dimensions()]
            import_path = root / "综合素质评价导入.xlsx"
            _write_import_workbook(import_path, dimensions)

            result = quality_controller.import_quality_from_excel(import_path, class_id)
            assert result.total_rows == 8
            assert result.matched_students == 2
            assert result.created_students == 0
            assert result.updated_evaluations == 12
            assert result.na_evaluations == 4
            assert not result.issues

            complete = quality_controller.get_student_evaluation(complete_student_id, SEMESTER_RULES[-1].key)
            assert complete["final_ready"]
            assert not complete["contains_na"]
            assert complete["final_score"] == 50
            assert complete["dimension_scores"] == {dimension: 10 for dimension in dimensions}

            transfer = quality_controller.get_student_evaluation(transfer_student_id, SEMESTER_RULES[-1].key)
            assert transfer["final_ready"]
            assert transfer["contains_na"]
            assert transfer["completed_terms"] == 6
            assert transfer["final_score"] == 25
            assert transfer["dimension_scores"] == {dimension: 5 for dimension in dimensions}
            assert transfer["na_semesters"] == [rule.key for rule in SEMESTER_RULES[:4]]
            first_term = quality_controller.get_student_evaluation(transfer_student_id, SEMESTER_RULES[0].key)
            assert first_term["ratings"] == {dimension: "N/A" for dimension in dimensions}

            export_path = quality_controller.export_quality_to_excel(root / "综合素质评价导出.xlsx", class_id)
            workbook = load_workbook(export_path, read_only=True)
            total_sheet = workbook["综合素质总表"]
            headers = [cell.value for cell in next(total_sheet.iter_rows(max_row=1))]
            rows = {
                row[0]: row
                for row in total_sheet.iter_rows(min_row=2, values_only=True)
            }
            workbook.close()
            assert rows["转入学生"][headers.index("初一上学期-思想品德")] == "N/A"
            assert rows["转入学生"][headers.index("完成状态")] == "含 N/A"

            standalone_result = quality_controller.import_quality_from_excel(import_path)
            assert standalone_result.matched_students == 2
            assert standalone_result.created_students == 2
            assert standalone_result.created_class
            assert standalone_result.target_class_name == "综合素质导入班"
            standalone_rows = quality_controller.list_student_overviews(standalone_result.target_class_id)
            assert len(standalone_rows) == 2
        finally:
            engine.dispose()


def _write_import_workbook(file_path: Path, dimensions: list[str]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, sheet_name in enumerate(("七上", "七下", "八上", "八下", "九上", "九下")):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["序号", "姓名", *dimensions])
        worksheet.append([1, "完整学生", *(["A"] * len(dimensions))])
        if index >= 4:
            worksheet.append([2, "转入学生", *(["A"] * len(dimensions))])
    workbook.create_sheet("总计")
    workbook.save(file_path)


if __name__ == "__main__":
    run()
    print("Quality Excel import smoke test passed.")
