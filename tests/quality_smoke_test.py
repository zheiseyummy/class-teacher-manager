from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def run() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        config.DATA_DIR = root / "data"
        config.BACKUP_DIR = root / "backups"
        config.DATABASE_PATH = config.DATA_DIR / "class_manager.db"
        config.DATABASE_URL = f"sqlite:///{config.DATABASE_PATH.as_posix()}"

        from controllers.quality_controller import QualityController
        from controllers.student_controller import StudentController
        from database.connection import engine
        from database.init_db import initialize_database
        from openpyxl import load_workbook
        from utils.quality_scoring import SEMESTER_RULES

        try:
            initialize_database()
            student_controller = StudentController()
            class_id = student_controller.create_class({"name": "初三1班", "grade": "初三"})
            student_id = student_controller.create_student(
                {"class_id": class_id, "name": "测试学生", "student_no": "1001"}, []
            )

            quality_controller = QualityController()
            dimensions = [item["name"] for item in quality_controller.list_dimensions()]
            assert dimensions == ["思想品德", "学业水平", "身心健康", "艺术素养", "实践与创新"]

            quality_controller.save_evaluation(
                student_id,
                SEMESTER_RULES[0].key,
                {dimension: "B" for dimension in dimensions},
            )
            first_term = quality_controller.get_student_evaluation(student_id, SEMESTER_RULES[0].key)
            assert first_term["total_score"] == 3.6
            assert not first_term["final_ready"]

            for rule in SEMESTER_RULES[1:]:
                quality_controller.save_evaluation(
                    student_id,
                    rule.key,
                    {dimension: "A" for dimension in dimensions},
                )
            final_term = quality_controller.get_student_evaluation(student_id, SEMESTER_RULES[-1].key)
            assert final_term["completed_terms"] == 6
            assert final_term["final_ready"]
            assert final_term["final_score"] == 49.6
            assert final_term["dimension_scores"] == {dimension: 9.92 for dimension in dimensions}
            assert final_term["completed_dimensions"] == dimensions

            overview = quality_controller.list_student_overviews()[0]
            assert overview["最终分数"] == "49.6"

            export_path = quality_controller.export_quality_to_excel(root / "综合素质评价.xlsx")
            workbook = load_workbook(export_path, read_only=True)
            assert workbook.sheetnames == ["综合素质总表", "五维得分", "学期得分", "最终评定"]
            dimension_sheet = workbook["五维得分"]
            headers = [cell.value for cell in next(dimension_sheet.iter_rows(max_row=1))]
            row = [cell.value for cell in next(dimension_sheet.iter_rows(min_row=2, max_row=2))]
            workbook.close()
            assert "思想品德得分" in headers and "五维总分" in headers
            assert row[headers.index("五维总分")] == 49.6
        finally:
            engine.dispose()


if __name__ == "__main__":
    run()
    print("Quality evaluation smoke test passed.")
