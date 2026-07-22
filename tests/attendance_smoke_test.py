from __future__ import annotations

import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def run() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        config.DATA_DIR = root / "data"
        config.BACKUP_DIR = root / "backups"
        config.DATABASE_PATH = config.DATA_DIR / "class_manager.db"
        config.DATABASE_URL = f"sqlite:///{config.DATABASE_PATH.as_posix()}"

        from controllers.attendance_controller import AttendanceController, AttendanceDataError
        from controllers.student_controller import StudentController
        from database.connection import engine
        from database.init_db import initialize_database
        from openpyxl import load_workbook

        try:
            initialize_database()
            students = StudentController()
            class_a = students.create_class({"name": "初三1班", "grade": "初三"})
            class_b = students.create_class({"name": "初三2班", "grade": "初三"})
            student_a = students.create_student(
                {"class_id": class_a, "name": "李明", "student_no": "001"}, []
            )
            student_b = students.create_student(
                {"class_id": class_b, "name": "张婷", "student_no": "002"}, []
            )

            controller = AttendanceController()
            leave_id = controller.create_record(
                {
                    "student_id": student_a,
                    "record_type": "病假",
                    "approval_status": "已批准",
                    "start_time": datetime(2026, 9, 1, 8, 0),
                    "end_time": datetime(2026, 9, 2, 10, 0),
                    "reason": "发热就医",
                    "note": "已与家长确认",
                }
            )
            late_id = controller.create_record(
                {
                    "student_id": student_a,
                    "record_type": "迟到",
                    "approval_status": "已登记",
                    "start_time": datetime(2026, 9, 5, 7, 50),
                    "end_time": datetime(2026, 9, 5, 8, 10),
                    "reason": "交通拥堵",
                }
            )
            controller.create_record(
                {
                    "student_id": student_b,
                    "record_type": "缺勤",
                    "approval_status": "待审批",
                    "start_time": datetime(2026, 9, 6, 8, 0),
                    "end_time": datetime(2026, 9, 6, 17, 0),
                    "reason": "未到校",
                }
            )

            class_a_rows = controller.list_records(class_id=class_a)
            assert len(class_a_rows) == 2
            assert {row["record_type"] for row in class_a_rows} == {"病假", "迟到"}
            summary = controller.get_summary(class_id=class_a)
            assert summary == {"total": 2, "leave": 1, "late": 1, "early_leave": 0, "absence": 0}

            leave = controller.get_record(leave_id)
            assert leave["duration_display"] == "1天2小时"
            assert controller.list_records(keyword="交通")[-1]["id"] == late_id
            assert len(controller.list_records(record_type="缺勤")) == 1
            assert len(
                controller.list_records(start_date=date(2026, 9, 5), end_date=date(2026, 9, 5))
            ) == 1

            controller.update_record(
                late_id,
                {
                    "student_id": student_a,
                    "record_type": "早退",
                    "approval_status": "已登记",
                    "start_time": datetime(2026, 9, 5, 16, 10),
                    "end_time": datetime(2026, 9, 5, 17, 0),
                    "reason": "身体不适",
                    "note": "家长已接回",
                },
            )
            updated = controller.get_record(late_id)
            assert updated["record_type"] == "早退"
            assert updated["duration_minutes"] == 50
            assert controller.get_summary(class_id=class_a)["early_leave"] == 1

            try:
                controller.create_record(
                    {
                        "student_id": student_a,
                        "record_type": "事假",
                        "approval_status": "已批准",
                        "start_time": datetime(2026, 9, 8, 10, 0),
                        "end_time": datetime(2026, 9, 8, 8, 0),
                    }
                )
            except AttendanceDataError:
                pass
            else:
                raise AssertionError("An invalid time range should be rejected.")

            export_path = controller.export_records_to_excel(root / "请假与考勤.xlsx", class_id=class_a)
            workbook = load_workbook(export_path, read_only=True)
            worksheet = workbook["请假与考勤"]
            headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
            rows = list(worksheet.iter_rows(min_row=2, values_only=True))
            workbook.close()
            assert headers == ["日期时间", "结束时间", "学生", "学号", "班级", "类型", "状态", "时长", "事由", "备注"]
            assert len(rows) == 2

            controller.delete_record(late_id)
            assert len(controller.list_records(class_id=class_a)) == 1
        finally:
            engine.dispose()


if __name__ == "__main__":
    run()
    print("Attendance management smoke test passed.")
